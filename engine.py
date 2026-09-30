"""Ядро «Желтка»: каталог, анкета, проверки, расчёты, сборка .docx.

Никакого ИИ и никакой сети: только шаблоны. Этот модуль не знает про Telegram,
поэтому его можно тестировать и переиспользовать (например, для сайта).
"""
from __future__ import annotations

import datetime as dt
import io
import os
import re
from typing import Any

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from jinja2 import Environment, StrictUndefined

NBSP = " "


# ----------------------------------------------------------------- helpers
def today() -> dt.date:
    """Сегодня. Для тестов можно подменить переменной ENGINE_TODAY=ДД.ММ.ГГГГ."""
    forced = os.environ.get("ENGINE_TODAY")
    return dt.datetime.strptime(forced, "%d.%m.%Y").date() if forced else dt.date.today()


def fmt_date(d: dt.date) -> str:
    return d.strftime("%d.%m.%Y")


def money(x: Any) -> str:
    x = float(x or 0)
    if abs(x - round(x)) < 0.005:
        return f"{int(round(x)):,}".replace(",", NBSP)
    return f"{x:,.2f}".replace(",", NBSP).replace(".", ",")


def days_since(d: dt.date) -> int:
    return (today() - d).days


def add_days(d: dt.date, n: int) -> dt.date:
    return d + dt.timedelta(days=n)


# Шкала НДФЛ для резидентов по основной налоговой базе (ст. 224 НК РФ): (порог, ставка) по годам.
NDFL_SCALE = {
    2025: [(2_400_000, 0.13), (5_000_000, 0.15), (20_000_000, 0.18), (50_000_000, 0.20), (None, 0.22)],
    2021: [(5_000_000, 0.13), (None, 0.15)],
    0: [(None, 0.13)],
}


def ndfl(income: Any, year: int) -> float:
    """НДФЛ с годового дохода по шкале того года. Вычет уменьшает доход, поэтому возврат =
    ndfl(доход) - ndfl(доход - вычет): считается по самым высоким из ставок человека."""
    scale = next(v for k, v in sorted(NDFL_SCALE.items(), reverse=True) if year >= k)
    left, low, tax = max(float(income or 0), 0.0), 0.0, 0.0
    for top, rate in scale:
        part = left if top is None else min(left, top - low)
        if part <= 0:
            break
        tax += part * rate
        left -= part
        low = top or low
    return round(tax, 2)


def plural(n: Any, one: str, few: str, many: str) -> str:
    """1 день, 2 дня, 5 дней."""
    n = abs(int(n))
    if n % 10 == 1 and n % 100 != 11:
        return one
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return few
    return many


def short(fio: str) -> str:
    parts = fio.split()
    if not parts:
        return ""
    return parts[0] + (" " + " ".join(p[0] + "." for p in parts[1:3]) if len(parts) > 1 else "")


def lines(text: str) -> list[str]:
    return [s.strip() for s in re.split(r"[;\n]+", text or "") if s.strip()]


def fmt_rate(r: Any) -> str:
    r = float(r)
    return str(int(r)) if r == int(r) else f"{r:g}".replace(".", ",")


def doc_list(docs: list[str], period: str | None) -> list[str]:
    period = period or "[укажите период]"
    names = {
        "dohod": f"Справку о доходах и суммах налога физического лица за {period}.",
        "zarabotok": "Справку о сумме заработка для расчёта пособий по временной нетрудоспособности, "
                     "по беременности и родам, ежемесячного пособия по уходу за ребёнком.",
        "std": "Сведения о трудовой деятельности (форма СТД-Р) либо заверенную копию трудовой книжки.",
        "dogovor": "Заверенную копию трудового договора и дополнительных соглашений к нему.",
        "prikazy": "Заверенные копии приказов о приёме на работу, переводах и увольнении.",
        "listki": f"Расчётные листки за {period}.",
    }
    return [names[d] for d in docs if d in names]


def deduction_docs(types: list[str], year: int) -> list[str]:
    """Документы для социального вычета. С расходов 2024 года основной документ — справка об оплате
    от организации (форма ФНС); для более ранних лет — договор, лицензия и платёжные документы."""
    new = year >= 2024
    out = ["Паспорт и ИНН; реквизиты своей карты для возврата."]
    if "med" in types or "exp" in types:
        if new:
            out.append("Справка об оплате медицинских услуг для налоговой — выдаёт клиника по заявлению. Для расходов с 2024 года это основной документ.")
        else:
            out.append("Договор с клиникой, её лицензия (реквизиты обычно есть в договоре) и справка об оплате медицинских услуг для налоговой.")
    if "drugs" in types:
        out.append("Рецепт или назначение врача и кассовые чеки на лекарства. Вычет за лекарства дают только через декларацию 3-НДФЛ.")
    if "edu" in types or "child" in types:
        if new:
            out.append("Справка об оплате образовательных услуг для налоговой — выдаёт учебное заведение. Для расходов с 2024 года это основной документ.")
        else:
            out.append("Договор на обучение, лицензия организации (если реквизитов нет в договоре) и платёжные документы.")
    if "child" in types:
        out.append("Обучение детей, брата или сестры учитывается только при очной форме обучения и возрасте до 24 лет.")
    if "fit" in types:
        if new:
            out.append("Справка об оплате физкультурно-оздоровительных услуг — выдаёт клуб. Клуб должен быть в перечне Минспорта.")
        else:
            out.append("Договор с клубом и платёжные документы; клуб должен быть в перечне Минспорта.")
    out.append("Если платили за супруга, детей, брата, сестру или родителей — документы о родстве или браке (свидетельство о рождении, о браке).")
    return out


# ----------------------------------------------------------------- jinja
ENV = Environment(undefined=StrictUndefined, trim_blocks=True, lstrip_blocks=True,
                  keep_trailing_newline=False, autoescape=False)
ENV.globals.update(date=fmt_date, money=money, days_since=days_since, add_days=add_days,
                   short=short, plural=plural, ndfl=ndfl, lines=lines, fmt_rate=fmt_rate, doc_list=doc_list,
                   deduction_docs=deduction_docs)


# ----------------------------------------------------------------- catalog
PERSON_SECTION = "Ваши данные"


class CatalogError(Exception):
    pass


class Catalog:
    def __init__(self, path: str):
        with open(path, encoding="utf-8") as f:
            raw = yaml.safe_load(f)
        self.categories = raw["categories"]
        self.labels: dict[str, str] = raw.get("labels", {})
        common = raw.get("common", {})
        self.docs: dict[str, dict] = {}
        for doc in raw["docs"]:
            qs = []
            for q in doc["questions"]:
                if "use" in q:
                    qs.extend(dict(x, section=PERSON_SECTION) for x in common[q["use"]])
                else:
                    qs.append(q)
            doc["questions"] = qs
            self._sections(doc)
            self._check(doc)
            self.docs[doc["id"]] = doc

    def _sections(self, doc: dict) -> None:
        """Каждому вопросу — раздел. Порядок разделов = порядок первых вопросов."""
        by_key = {}
        for s in doc.get("sections") or []:
            for k in s["keys"]:
                by_key[k] = s["title"]
        first = (doc.get("sections") or [{"title": "Ситуация"}])[0]["title"]
        order: list[str] = []
        for q in doc["questions"]:
            q["section"] = q.get("section") or by_key.get(q["key"]) or (PERSON_SECTION if q["key"] == "fio" else first)
            q["label"] = q.get("label") or self.labels.get(q["key"], q["key"])
            if q["section"] not in order:
                order.append(q["section"])
        doc["_sections"] = order

    def _check(self, doc: dict) -> None:
        keys = [q["key"] for q in doc["questions"]]
        if len(keys) != len(set(keys)):
            raise CatalogError(f"{doc['id']}: повторяются ключи вопросов")
        for q in doc["questions"]:
            if q["type"] in ("choice", "multi") and not q.get("options"):
                raise CatalogError(f"{doc['id']}.{q['key']}: нет вариантов ответа")
        for fld in ("id", "category", "title", "button", "price", "template", "filename"):
            if fld not in doc:
                raise CatalogError(f"{doc.get('id')}: нет поля {fld}")

    def in_category(self, cat_id: str) -> list[dict]:
        return [d for d in self.docs.values() if d["category"] == cat_id]


# ----------------------------------------------------------------- answers
def parse_answer(q: dict, text: str) -> tuple[Any, str | None]:
    """Проверяет текстовый ответ. Возвращает (значение, ошибка)."""
    t = (text or "").strip()
    typ = q["type"]
    if q.get("optional") and t in ("-", "—", "нет", "Нет"):
        return "", None
    if not t:
        return None, "Ответ пустой. Напишите его, пожалуйста, текстом."
    if typ == "text":
        return (t, None) if len(t) <= 500 else (None, "Получилось длинно — сократите, пожалуйста, до 500 символов.")
    if typ == "long":
        return (t.rstrip("."), None) if len(t) <= 1500 else (None, "Получилось длинно — сократите, пожалуйста, до 1500 символов.")
    if typ == "fio":
        words = t.split()
        if len(words) < 2 or not all(re.fullmatch(r"[А-ЯЁа-яёA-Za-z\-]+", w) for w in words):
            return None, "Укажите фамилию и имя (и отчество, если есть) буквами, например: Иванов Иван Иванович."
        return " ".join(w[:1].upper() + w[1:] for w in words), None
    if typ == "money":
        s = re.sub(r"(руб\.?|р\.|₽|\s)", "", t, flags=re.I).replace(",", ".")
        try:
            v = float(s)
        except ValueError:
            return None, "Укажите сумму цифрами, например 24990 или 24990,50."
        if v <= 0 or v > 1e9:
            return None, "Сумма должна быть больше нуля."
        return round(v, 2), None
    if typ == "int":
        if not re.fullmatch(r"\d{1,6}", t):
            return None, "Укажите целое число цифрами, например 10."
        v = int(t)
        return (v, None) if v > 0 else (None, "Число должно быть больше нуля.")
    if typ == "date":
        m = re.fullmatch(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})", t)
        if not m:
            return None, "Укажите дату в формате ДД.ММ.ГГГГ, например 05.09.2026."
        try:
            d = dt.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            return None, "Такой даты нет — проверьте число и месяц."
        if q.get("past") and d > today():
            return None, "Эта дата ещё не наступила — проверьте год и месяц."
        if d < today() - dt.timedelta(days=365 * 30):
            return None, "Дата слишком давняя — проверьте год."
        return d, None
    return None, "Для этого вопроса выберите вариант кнопкой под сообщением."


def option_label(q: dict, v: Any) -> str:
    for o in q.get("options", []):
        if o["v"] == v:
            return o["label"]
    return str(v)


# ----------------------------------------------------------------- flow
def base_context(doc: dict, answers: dict, cfg: dict) -> dict:
    ctx: dict[str, Any] = {q["key"]: None for q in doc["questions"]}
    ctx.update(answers)
    for q in doc["questions"]:
        k = q["key"]
        if q["type"] == "choice":
            ctx[k + "_label"] = option_label(q, answers.get(k)) if k in answers else None
        if q["type"] == "multi":
            ctx[k] = answers.get(k) or []
            ctx[k + "_labels"] = [option_label(q, v) for v in ctx[k]]
    ctx.update(cfg)
    ctx["today"] = today()
    return ctx


def is_asked(q: dict, ctx: dict) -> bool:
    cond = q.get("when")
    return True if not cond else bool(ENV.compile_expression(cond)(**ctx))


def next_question(doc: dict, answers: dict, cfg: dict) -> dict | None:
    ctx = base_context(doc, answers, cfg)
    for q in doc["questions"]:
        if q["key"] in answers:
            continue
        if is_asked(q, ctx):
            return q
    return None


def prune(doc: dict, answers: dict, cfg: dict) -> dict:
    """Убирает ответы на вопросы, которые после исправления стали неактуальны
    (например, «куда вернуть деньги», если требование сменили на замену)."""
    answers = dict(answers)
    changed = True
    while changed:
        changed = False
        ctx = base_context(doc, answers, cfg)
        for q in doc["questions"]:
            if q["key"] in answers and not is_asked(q, ctx):
                del answers[q["key"]]
                changed = True
                break
    return answers


def section_info(doc: dict, key: str) -> tuple[int, int, str]:
    """(номер раздела с 1, всего разделов, название) для вопроса."""
    title = next(q["section"] for q in doc["questions"] if q["key"] == key)
    order = doc["_sections"]
    return order.index(title) + 1, len(order), title


def section_first_key(doc: dict, title: str) -> str:
    return next(q["key"] for q in doc["questions"] if q["section"] == title)


def fmt_value(q: dict, v: Any) -> str:
    if v is None or v == "" or v == []:
        return "не указано"
    t = q["type"]
    if t == "date":
        return fmt_date(v)
    if t == "money":
        return f"{money(v)} ₽"
    if t == "yesno":
        return "да" if v else "нет"
    if t == "choice":
        return option_label(q, v)
    if t == "multi":
        return ", ".join(option_label(q, x) for x in v)
    s = str(v)
    return s if len(s) <= 160 else s[:157] + "…"


def summary(doc: dict, answers: dict, cfg: dict) -> list[tuple[str, list[tuple[str, str]]]]:
    """Сводка для проверки: [(раздел, [(подпись, значение), ...]), ...].
    Показываются только вопросы, которые реально задавались в этом сценарии."""
    ctx = base_context(doc, answers, cfg)
    out: dict[str, list[tuple[str, str]]] = {t: [] for t in doc["_sections"]}
    for q in doc["questions"]:
        if q["key"] in answers or is_asked(q, ctx):
            out[q["section"]].append((q["label"], fmt_value(q, answers.get(q["key"]))))
    return [(t, rows) for t, rows in out.items() if rows]


def full_context(doc: dict, answers: dict, cfg: dict) -> dict:
    ctx = base_context(doc, answers, cfg)
    for k, expr in (doc.get("computed") or {}).items():
        ctx[k] = ENV.compile_expression(expr)(**ctx)
    for k, tpl in (doc.get("vars") or {}).items():
        ctx[k] = ENV.from_string(tpl).render(**ctx)
    return ctx


def check_warnings(doc: dict, answers: dict, cfg: dict) -> list[dict]:
    ctx = full_context(doc, answers, cfg)
    out = []
    for w in doc.get("warnings") or []:
        if ENV.compile_expression(w["when"])(**ctx):
            out.append({"text": w["text"], "block": bool(w.get("block"))})
    return out


def render_text(doc: dict, answers: dict, cfg: dict) -> str:
    ctx = full_context(doc, answers, cfg)
    return ENV.from_string(doc["template"]).render(**ctx)


def render_snippet(tpl: str, doc: dict, answers: dict, cfg: dict) -> str:
    return ENV.from_string(tpl).render(**full_context(doc, answers, cfg)).strip()


# ----------------------------------------------------------------- docx
def _clean_lines(text: str) -> list[str]:
    out: list[str] = []
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            if out and out[-1] != "":
                out.append("")
            continue
        out.append(line.strip())
    while out and out[-1] == "":
        out.pop()
    return out


def _font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Times New Roman")


def _page_number_footer(section) -> None:
    """Номер страницы внизу по центру (поле PAGE) — мелко, без брендинга."""
    from docx.oxml import OxmlElement
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _font(run, size=9)
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "end"})):
        el = OxmlElement(tag)
        if attr:
            for k, v in attr.items():
                el.set(qn(k), v)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        run._r.append(el)


def _clean_properties(d, title: str = "") -> None:
    cp = d.core_properties
    cp.author, cp.last_modified_by, cp.comments, cp.keywords = "", "", "", ""
    cp.title = title


def build_memo(title: str, steps_html: str, logo_path: str, bot_link: str = "") -> bytes:
    """Фирменная памятка «что делать дальше» — отдельный файл, который не отправляют адресату."""
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(11)
    st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(4), 1.25

    def run(p, text, size=11, bold=False, color=None):
        r = p.add_run(text)
        r.font.name, r.font.size, r.font.bold = "Arial", Pt(size), bold
        rpr = r._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Arial")
        if color:
            from docx.shared import RGBColor
            r.font.color.rgb = RGBColor.from_string(color)
        return r

    if logo_path and os.path.exists(logo_path):
        d.add_picture(logo_path, width=Cm(4.2))
    p = d.add_paragraph()
    run(p, "Документы. По шагам.", 10, color="64665D")
    # тонкая жёлтая линия-акцент
    from docx.oxml import OxmlElement
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "6"), ("w:color", "F4C542")):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)

    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    run(p, "Памятка", 10, color="64665D")
    p = d.add_paragraph()
    run(p, title, 18, bold=True, color="24251F")
    p = d.add_paragraph()
    run(p, "Эта памятка — для вас. Адресату отправляйте только сам документ.", 10, color="64665D")

    for raw in steps_html.split("\n"):
        line = raw.strip()
        if not line:
            continue
        bold = line.startswith("<b>") and line.endswith("</b>")
        text = re.sub(r"<[^>]+>", "", line)
        p = d.add_paragraph()
        if bold:
            p.paragraph_format.space_before = Pt(12)
            run(p, text, 13, bold=True, color="24251F")
        else:
            run(p, text, 11, color="24251F")

    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    run(p, "Что важно помнить", 13, bold=True, color="24251F")
    for t in ("Сохраните копию документа с отметкой о получении или квитанцию об отправке — это главное доказательство.",
              "Отметьте отправку в боте — он напомнит, когда проверять ответ, и подскажет следующий шаг.",
              "Документ собран по шаблону и не является юридической консультацией."):
        p = d.add_paragraph()
        run(p, "— " + t, 11, color="24251F")

    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    run(p, "желток", 11, bold=True, color="24251F")
    run(p, "  ·  Документы. По шагам." + (f"  ·  {bot_link}" if bot_link else ""), 10, color="64665D")
    _clean_properties(d, f"Памятка — {title}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_docx(text: str, title: str = "") -> bytes:
    d = Document()
    _clean_properties(d, title)
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    text_width = sec.page_width - sec.left_margin - sec.right_margin
    _page_number_footer(sec)

    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    gap_next = False
    for line in _clean_lines(text):
        if line == "":
            gap_next = True
            continue
        p = d.add_paragraph()
        pf = p.paragraph_format
        if gap_next:
            pf.space_before = Pt(8)
            gap_next = False
        if line.startswith(">"):
            body = line[1:].strip()
            pf.left_indent = Cm(9.0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _font(p.add_run(body))
        elif line.startswith("### "):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.space_before = Pt(10)
            pf.space_after = Pt(3)
            pf.keep_with_next = True
            _font(p.add_run(line[4:].strip()), bold=True)
        elif line.startswith("## "):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_after = Pt(6)
            _font(p.add_run(line[3:].strip()))
        elif line.startswith("# "):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(14)
            _font(p.add_run(line[2:].strip()), size=14, bold=True)
        elif line.startswith("~ "):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _font(p.add_run(line[2:].strip()))
        elif line.startswith("= "):
            left, _, right = line[2:].partition(" | ")
            pf.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
            pf.space_before = Pt(18)
            _font(p.add_run(f"{left.strip()}\t{right.strip()}"))
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.first_line_indent = Cm(1.25)
            pf.space_after = Pt(4)
            _font(p.add_run(line))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def preview_text(text: str, share: float = 0.45) -> tuple[str, int]:
    """Первая часть документа для предпросмотра до оплаты (без разметки)."""
    ls = [l for l in _clean_lines(text) if l and not l.startswith(">")]   # шапку «Кому/От» пропускаем
    cut = max(4, int(len(ls) * share))
    shown = []
    for l in ls[:cut]:
        for pre in ("### ", "## ", "# ", "~ ", "= ", ">"):
            if l.startswith(pre):
                l = l[len(pre):].strip()
                break
        shown.append(l.replace(" | ", "   "))
    return "\n".join(s for s in shown if s), max(0, len(ls) - cut)
